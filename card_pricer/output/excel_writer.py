import os
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from engine.classify import Card
from config import OUTPUT_DIR

# ── Styles ─────────────────────────────────────────────────────────────────────
HDR_FONT      = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HDR_FILL      = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
BODY_FONT     = Font(name="Arial", size=10)
LINK_FONT     = Font(name="Arial", size=10, color="0563C1", underline="single")
MISSING_FILL  = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
SUMMARY_FILL  = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
SUMMARY_FILL2 = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
CENTER        = Alignment(horizontal="center", vertical="center")
WRAP          = Alignment(horizontal="left",   vertical="center")
THIN          = Border(
    bottom=Side(style="thin",   color="D9D9D9"),
    right =Side(style="thin",   color="D9D9D9"),
)
THICK_BOTTOM  = Border(bottom=Side(style="medium", color="1F4E79"))

# ── Column definitions ─────────────────────────────────────────────────────────
COLUMNS = [
    # Card identity
    ("Card Name",             32, WRAP,   "name"),
    ("Set",                   22, WRAP,   "set_name"),
    ("Card #",                10, CENTER, "number"),
    ("Condition",             11, CENTER, "condition_or_grade"),
    ("Grading",                9, CENTER, "grading_company"),
    ("Foil",                  11, CENTER, "foil"),
    ("Qty",                    5, CENTER, "quantity"),
    # Prices
    ("Collectr Price",        14, CENTER, "collectr_price"),
    ("eBay Listed",           12, CENTER, "ebay_listing_price"),  # from Collectr JSON
    ("TCGplayer Market",      16, CENTER, "tcgplayer_price"),
    ("PriceCharting Price",   18, CENTER, "pricecharting_price"),
    ("Price Source",          24, WRAP,   "price_source_used"),
    # Verification links
    ("Collectr Listing",      18, CENTER, "_link_collectr"),
    ("TCGplayer Search",      18, CENTER, "_link_tcg"),
    ("PriceCharting Page",    20, CENTER, "_link_pc"),
]

TOTAL_COLS = len(COLUMNS)
PRICE_KEYS = {"collectr_price", "ebay_listing_price",
              "tcgplayer_price", "pricecharting_price"}


def _make_hyperlink(ws, row: int, col: int, url: str, label: str):
    if not url:
        c = ws.cell(row=row, column=col, value="—")
        c.font = BODY_FONT
        c.alignment = CENTER
        return
    c = ws.cell(row=row, column=col, value=label)
    c.hyperlink = url
    c.font = LINK_FONT
    c.alignment = CENTER


def _write_summary_rows(ws, cards: list[Card]):
    collectr_total = sum(
        (c.collectr_price or 0) * (c.quantity or 1) for c in cards
    )
    market_total = 0.0
    unpriced = 0
    for c in cards:
        qty  = c.quantity or 1
        best = c.tcgplayer_price or c.pricecharting_price or c.collectr_price
        if best:
            market_total += best * qty
        else:
            unpriced += 1

    # Row 1 — Collectr value
    ws.row_dimensions[1].height = 22
    l1 = ws.cell(1, 1, "Collectr Collection Value")
    l1.font = Font(name="Arial", bold=True, size=11)
    l1.fill = SUMMARY_FILL
    l1.alignment = WRAP
    v1 = ws.cell(1, 2, collectr_total)
    v1.font = Font(name="Arial", bold=True, size=13, color="375623")
    v1.number_format = '"$"#,##0.00'
    v1.fill = SUMMARY_FILL
    v1.alignment = CENTER
    n1 = ws.cell(1, 3, f"Sum of Collectr prices × qty for all {len(cards)} cards")
    n1.font = Font(name="Arial", size=9, italic=True, color="666666")
    n1.fill = SUMMARY_FILL
    n1.alignment = WRAP
    for col in range(4, TOTAL_COLS + 1):
        ws.cell(1, col, "").fill = SUMMARY_FILL

    # Row 2 — Market value
    ws.row_dimensions[2].height = 22
    l2 = ws.cell(2, 1, "Market Collection Value")
    l2.font = Font(name="Arial", bold=True, size=11)
    l2.fill = SUMMARY_FILL2
    l2.alignment = WRAP
    v2 = ws.cell(2, 2, market_total)
    v2.font = Font(name="Arial", bold=True, size=13, color="1F3864")
    v2.number_format = '"$"#,##0.00'
    v2.fill = SUMMARY_FILL2
    v2.alignment = CENTER
    unpriced_note = f" ({unpriced} cards unpriced)" if unpriced else ""
    n2 = ws.cell(
        2, 3,
        f"Best price × qty — TCGplayer → PriceCharting → Collectr{unpriced_note}"
    )
    n2.font = Font(name="Arial", size=9, italic=True, color="666666")
    n2.fill = SUMMARY_FILL2
    n2.alignment = WRAP
    for col in range(4, TOTAL_COLS + 1):
        ws.cell(2, col, "").fill = SUMMARY_FILL2
    for col in range(1, TOTAL_COLS + 1):
        ws.cell(2, col).border = THICK_BOTTOM


def write_cards_to_excel(cards: list[Card], filename: str | None = None) -> str:
    if filename is None:
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"card_pricing_results_{ts}.xlsx"

    filepath = os.path.join(OUTPUT_DIR, filename)
    wb = Workbook()
    ws = wb.active
    ws.title = "Pricing Results"

    HEADER_ROW = 3
    DATA_START  = 4

    _write_summary_rows(ws, cards)

    # Header row
    ws.row_dimensions[HEADER_ROW].height = 18
    for col_idx, (header, width, align, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col_idx, value=header)
        cell.font = HDR_FILL and HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = CENTER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Data rows
    for row_offset, card in enumerate(cards):
        row_idx = DATA_START + row_offset
        ws.row_dimensions[row_idx].height = 16

        for col_idx, (_, _, align, key) in enumerate(COLUMNS, start=1):
            if key.startswith("_link_"):
                if key == "_link_collectr":
                    _make_hyperlink(ws, row_idx, col_idx,
                                    card.collectr_item_url, "View on Collectr")
                elif key == "_link_tcg":
                    _make_hyperlink(ws, row_idx, col_idx,
                                    card.tcgplayer_search_url, "TCGplayer Search")
                elif key == "_link_pc":
                    _make_hyperlink(ws, row_idx, col_idx,
                                    card.pricecharting_result_url, "PriceCharting")
                continue

            if key == "grading_company":
                # Show grade info even when company is unknown
                if card.card_type == "graded":
                    grade_display = ""
                    if card.grading_company:
                        grade_display = f"{card.grading_company} {card.grade_value or ''}".strip()
                    elif card.collectr_grade_value:
                        grade_display = f"Grade {card.collectr_grade_value}"
                    elif card.condition_or_grade:
                        grade_display = card.condition_or_grade
                    cell = ws.cell(row=row_idx, column=col_idx, value=grade_display)
                else:
                    cell = ws.cell(row=row_idx, column=col_idx, value="")
                cell.font = BODY_FONT
                cell.alignment = CENTER
                cell.border = THIN
                continue

            value = getattr(card, key, None)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.alignment = align
            cell.border = THIN

            if key in PRICE_KEYS:
                cell.alignment = CENTER
                if value is not None:
                    cell.number_format = '"$"#,##0.00'
                else:
                    cell.fill = MISSING_FILL

    ws.freeze_panes = f"A{DATA_START}"
    ws.auto_filter.ref = (
        f"A{HEADER_ROW}:{get_column_letter(TOTAL_COLS)}{HEADER_ROW}"
    )
    wb.save(filepath)
    return filepath
